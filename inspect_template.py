import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = pptx.Presentation(r'C:\Users\saide\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
print(f"Presentation dimensions: {prs.slide_width/914400:.2f} in x {prs.slide_height/914400:.2f} in")
print(f"Total slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides, 1):
    print(f"\n==================== SLIDE {i} ====================")
    for s in slide.shapes:
        t = ""
        if s.has_text_frame:
            t = " // ".join(p.text.strip() for p in s.text_frame.paragraphs if p.text.strip())
        print(f"ID={s.shape_id:5d} | Name={s.name:30s} | Type={str(s.shape_type):20s} | Pos=({s.left/914400:5.2f}, {s.top/914400:5.2f}, {s.width/914400:5.2f}, {s.height/914400:5.2f}) | Text='{t[:80]}'")
